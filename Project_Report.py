from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
doc = Document()
# =========================================================
# TITLE PAGE
# =========================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("INTERNSHIP SKILL-GAP ANALYTICS SYSTEM")
run.bold = True
run.font.size = Pt(20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("A Data Analytics Project")
run.italic = True
run.font.size = Pt(14)
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Submitted by:")
run.bold = True
run.font.size = Pt(14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Komal Kumar")
run.bold = True
run.font.size = Pt(16)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(
    "Course: B.Tech – Computer Science & Engineering"
).font.size = Pt(13)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(
    "Technology: Python | Pandas | CSV | Power BI"
).font.size = Pt(13)
# =========================================================
# INTRODUCTION
# =========================================================
doc.add_page_break()
doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "The Internship Skill-Gap Analytics System is a data analytics project "
    "developed to analyze the skills possessed by students and compare them "
    "with the skills required for internship and job roles."
)
doc.add_paragraph(
    "The system helps students understand their current skill level, identify "
    "missing skills and evaluate their readiness for a particular internship "
    "or job role. Python and Pandas are used for data processing, while "
    "Power BI is used to create an interactive dashboard."
)
# =========================================================
# PROBLEM STATEMENT
# =========================================================
doc.add_heading("2. Problem Statement", level=1)
doc.add_paragraph(
    "Students often face difficulty in identifying whether their existing "
    "technical skills match the requirements of internship and job roles."
)
doc.add_paragraph(
    "The main problem is the lack of a simple data-driven system that can "
    "compare student skills with required skills and clearly show the missing "
    "skills. This project addresses this problem by providing skill matching, "
    "missing skill identification and internship readiness analysis."
)
# =========================================================
# OBJECTIVES
# =========================================================
doc.add_heading("3. Objectives", level=1)
objectives = [
    "To analyze the technical skills of students.",
    "To compare student skills with internship or job requirements.",
    "To calculate the Skill Match Percentage.",
    "To identify missing or required skills.",
    "To evaluate internship readiness.",
    "To provide meaningful insights through Power BI visualizations.",
    "To help students understand areas where skill improvement is required."
]
for item in objectives:
    doc.add_paragraph(item, style="List Bullet")
# =========================================================
# TECHNOLOGIES USED
# =========================================================
doc.add_heading("4. Technologies Used", level=1)
technologies = [
    ("Python", "Used for data processing and analysis."),
    ("Pandas", "Used for reading, cleaning and analyzing data."),
    ("CSV", "Used for storing student and skill data."),
    ("Power BI", "Used for dashboard creation and data visualization."),
    ("Microsoft Excel", "Used for preparing and managing the dataset.")
]
for tech, use in technologies:
    p = doc.add_paragraph()
    r = p.add_run(tech + ": ")
    r.bold = True
    p.add_run(use)
# =========================================================
# METHODOLOGY
# =========================================================
doc.add_heading("5. Methodology", level=1)
steps = [
    "Collect student skill and internship requirement data.",
    "Store the collected data in CSV format.",
    "Load the data using Python and Pandas.",
    "Clean and prepare the dataset.",
    "Compare student skills with required skills.",
    "Calculate Skill Match Percentage.",
    "Identify missing skills for each student.",
    "Analyze internship readiness.",
    "Import the processed data into Power BI.",
    "Create an interactive dashboard."
]
for i, step in enumerate(steps, 1):
    doc.add_paragraph(f"{i}. {step}")
# =========================================================
# DATASET
# =========================================================
doc.add_heading("6. Dataset Description", level=1)
doc.add_paragraph(
    "The dataset contains information related to students, their technical "
    "skills and the skills required for different internship or job roles."
)
dataset_fields = [
    "Student ID",
    "Student Name",
    "Target Role",
    "Student Skills",
    "Required Skills",
    "Skill Match Percentage",
    "Missing Skills",
    "Readiness Status"
]
for field in dataset_fields:
    doc.add_paragraph(field, style="List Bullet")
# =========================================================
# DATA PROCESSING
# =========================================================
doc.add_heading("7. Data Processing Using Python", level=1)

doc.add_paragraph(
    "Python and Pandas are used to load and process the dataset. The data is "
    "checked for missing values, duplicate records and inconsistent entries."
)
doc.add_paragraph(
    "After cleaning the data, student skills are compared with the required "
    "skills of the selected internship or job role. The processed information "
    "is then used for further analysis and visualization."
)
# =========================================================
# SKILL GAP ANALYSIS
# =========================================================
doc.add_heading("8. Skill-Gap Analysis", level=1)
doc.add_paragraph(
    "Skill-gap analysis determines the difference between the skills a student "
    "currently possesses and the skills required for a target internship or "
    "job role."
)
doc.add_paragraph(
    "The analysis provides three important outputs: Skill Match Percentage, "
    "Missing Skills and Internship Readiness."
)
# =========================================================
# POWER BI DASHBOARD
# =========================================================
doc.add_heading("9. Power BI Dashboard", level=1)
doc.add_paragraph(
    "Power BI is used to create an interactive dashboard for presenting the "
    "results of the skill-gap analysis."
)
dashboard_items = [
    "Total Students",
    "Skill Match Percentage",
    "Missing Skills",
    "Internship Readiness",
    "Target Role Analysis",
    "Student Skill Analysis"
]
for item in dashboard_items:
    doc.add_paragraph(item, style="List Bullet")
# =========================================================
# RESULTS
# =========================================================
doc.add_heading("10. Results and Analysis", level=1)
doc.add_paragraph(
    "The system provides a clear view of student skill profiles and their "
    "compatibility with internship requirements. The dashboard helps identify "
    "students with stronger skill matches and highlights the skills that need "
    "improvement."
)
doc.add_paragraph(
    "The analysis can be used by students to understand their skill gaps and "
    "focus on learning the technologies required for their target roles."
)
# =========================================================
# ADVANTAGES
# =========================================================
doc.add_heading("11. Advantages", level=1)
advantages = [
    "Easy identification of missing skills.",
    "Data-driven internship readiness analysis.",
    "Interactive Power BI dashboard.",
    "Helps students plan their skill development.",
    "Reduces manual comparison of skills.",
    "Provides clear and understandable visual insights."
]
for item in advantages:
    doc.add_paragraph(item, style="List Bullet")
# =========================================================
# FUTURE SCOPE
# =========================================================
doc.add_heading("12. Future Scope", level=1)
future = [
    "Integration with real-time job and internship data.",
    "Automatic skill extraction from student resumes.",
    "Machine Learning based internship recommendation.",
    "Personalized learning recommendations.",
    "Web-based version of the system.",
    "Integration with online learning platforms."
]

for item in future:
    doc.add_paragraph(item, style="List Bullet")


# =========================================================
# CONCLUSION
# =========================================================

doc.add_heading("13. Conclusion", level=1)

doc.add_paragraph(
    "The Internship Skill-Gap Analytics System provides a practical approach "
    "to understanding the difference between student skills and internship "
    "requirements. By combining Python, Pandas, CSV data and Power BI, the "
    "system converts raw skill information into useful analytical insights."
)

doc.add_paragraph(
    "The project demonstrates how data analytics can be used to support "
    "students in identifying skill gaps and improving their preparation for "
    "internship and job opportunities."
)


# =========================================================
# SAVE DOCUMENT
# =========================================================

doc.save("Project_Report.docx")

print("Complete Project Report created successfully!")